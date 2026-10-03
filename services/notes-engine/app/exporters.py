from __future__ import annotations

import io
import re
from datetime import datetime, timezone
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Table, TableStyle
from reportlab.lib import colors
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from .consolidator import ConsolidationResult


def safe_filename(value: str) -> str:
    value = re.sub(r"[^A-Za-z0-9._-]+", "-", value).strip("-")
    return value[:100] or "lecture-notes"


def build_docx(unit_code: str, unit_name: str, topic: str, result: ConsolidationResult) -> bytes:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Inches(0.65)
    sec.left_margin = sec.right_margin = Inches(0.75)
    doc.styles["Normal"].font.name = "Aptos"
    doc.styles["Normal"].font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(f"{unit_name}\nUNIFIED LECTURE NOTES")
    run.bold = True
    run.font.size = Pt(20)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(f"{unit_code} · {topic}").italic = True

    doc.add_heading("Source Coverage", level=1)
    doc.add_paragraph(
        f"This document consolidates {len(result.source_materials)} complete source material(s). "
        f"Approximately {result.source_word_count:,} source words were read. "
        f"{result.duplicate_paragraph_count:,} exact duplicate passages were removed; materially different content was retained."
    )
    for source in result.source_materials:
        doc.add_paragraph(source, style="List Bullet")

    for section in result.sections:
        doc.add_heading(section.heading, level=1)
        for paragraph in section.paragraphs:
            text = (paragraph.replace("\\times", "×").replace("\\geq", "≥")
                    .replace("\\leq", "≤").replace("\\rightarrow", "→")
                    .replace("\\approx", "≈").replace("$$", "").replace("$", ""))
            if text.startswith(("- ", "* ")):
                doc.add_paragraph(text[2:], style="List Bullet")
            else:
                doc.add_paragraph(text)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def build_pdf(unit_code: str, unit_name: str, topic: str, result: ConsolidationResult) -> bytes:
    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=A4, rightMargin=16*mm, leftMargin=16*mm, topMargin=16*mm, bottomMargin=16*mm)
    styles = getSampleStyleSheet()
    title = ParagraphStyle("Title2", parent=styles["Title"], fontSize=19, leading=23, alignment=1, spaceAfter=10)
    subtitle = ParagraphStyle("Subtitle2", parent=styles["Normal"], fontSize=9.5, leading=12, alignment=1, spaceAfter=14)
    h1 = ParagraphStyle("H1", parent=styles["Heading1"], fontSize=14, leading=17, spaceBefore=12, spaceAfter=6)
    body = ParagraphStyle("Body2", parent=styles["BodyText"], fontSize=9.4, leading=13, spaceAfter=6)
    bullet = ParagraphStyle("Bullet2", parent=body, leftIndent=12, firstLineIndent=-6)
    story = [Paragraph(f"{unit_name}<br/>UNIFIED LECTURE NOTES", title), Paragraph(f"{unit_code} · {topic}", subtitle)]
    story.append(Paragraph("Source Coverage", h1))
    story.append(Paragraph(
        f"This document consolidates {len(result.source_materials)} complete source material(s). "
        f"Approximately {result.source_word_count:,} source words were read. "
        f"{result.duplicate_paragraph_count:,} exact duplicate passages were removed; materially different content was retained.", body))
    for source in result.source_materials:
        story.append(Paragraph(f"• {source}", bullet))
    for section in result.sections:
        story.append(Paragraph(section.heading, h1))
        for paragraph in section.paragraphs:
            text = (paragraph.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))
            text = (text.replace("\\times", "×").replace("\\geq", "≥")
                    .replace("\\leq", "≤").replace("\\rightarrow", "→")
                    .replace("\\approx", "≈").replace("$$", "").replace("$", ""))
            if text.startswith(("- ", "* ")):
                story.append(Paragraph(f"• {text[2:]}", bullet))
            else:
                story.append(Paragraph(text, body))
    doc.build(story)
    return out.getvalue()
