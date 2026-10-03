from __future__ import annotations

import io
import os
import zipfile
from pathlib import PurePosixPath

from docx import Document
from pypdf import PdfReader

MAX_ZIP_ENTRIES = 250
MAX_ZIP_MEMBER_BYTES = 30 * 1024 * 1024
MAX_ZIP_TOTAL_BYTES = 150 * 1024 * 1024
SUPPORTED = {".pdf", ".docx", ".txt", ".md"}


def clean(value: str) -> str:
    return value.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n").strip()


def extract_docx(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts: list[str] = []
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text.strip())
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " / ") for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    pages: list[str] = []
    for page in reader.pages:
        text = page.extract_text() or ""
        if text.strip():
            pages.append(text)
    return "\n".join(pages)


def safe_zip_name(name: str) -> str | None:
    normalized = name.replace("\\", "/").lstrip("/")
    if not normalized or normalized.endswith("/"):
        return None
    parts = PurePosixPath(normalized).parts
    if ".." in parts or normalized.startswith("__MACOSX/"):
        return None
    return normalized


def extract_zip(data: bytes) -> list[tuple[str, str]]:
    results: list[tuple[str, str]] = []
    total = 0
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        infos = [i for i in archive.infolist() if not i.is_dir()]
        if len(infos) > MAX_ZIP_ENTRIES:
            raise ValueError(f"ZIP contains too many files ({len(infos)}). Maximum is {MAX_ZIP_ENTRIES}.")
        for info in infos:
            name = safe_zip_name(info.filename)
            if not name:
                continue
            suffix = os.path.splitext(name)[1].lower()
            if suffix not in SUPPORTED:
                continue
            if info.file_size > MAX_ZIP_MEMBER_BYTES:
                raise ValueError(f"ZIP member is too large: {name}")
            total += info.file_size
            if total > MAX_ZIP_TOTAL_BYTES:
                raise ValueError("The uncompressed ZIP contents exceed the allowed total size.")
            member = archive.read(info)
            if suffix == ".pdf":
                text = extract_pdf(member)
            elif suffix == ".docx":
                text = extract_docx(member)
            else:
                text = member.decode("utf-8", errors="replace")
            text = clean(text)
            if text:
                results.append((name, text))
    if not results:
        raise ValueError("The ZIP contains no readable PDF, DOCX, TXT or Markdown sources.")
    return results


def extract_material(title: str, source_type: str, data: bytes | None, content_text: str | None) -> list[tuple[str, str]]:
    if content_text and content_text.strip():
        return [(title, content_text)]
    if not data:
        raise ValueError(f"No source data is available for {title}.")
    if source_type == "pdf":
        return [(title, extract_pdf(data))]
    if source_type == "docx":
        return [(title, extract_docx(data))]
    if source_type == "zip":
        members = extract_zip(data)
        return [(f"{title} / {name}", text) for name, text in members]
    if source_type == "text":
        return [(title, data.decode("utf-8", errors="replace"))]
    raise ValueError(f"Unsupported source type: {source_type}")
